from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor


REPORTS_DIR = Path("storage/reports")


BLOCK_ORDER = [
    "base",
    "present",
    "past",
    "future",
    "people",
    "project",
    "unusual",
]


BLOCK_TITLES = {
    "base": "БАЗА",
    "present": "НАСТОЯЩЕЕ ВРЕМЯ",
    "past": "ПРОШЛОЕ ВРЕМЯ",
    "future": "БУДУЩЕЕ ВРЕМЯ",
    "people": "ЛЮДИ И ОТНОШЕНИЯ",
    "project": "ПРОЕКТЫ, ДОСТИЖЕНИЯ, БИЗНЕС",
    "unusual": "НЕОБЫЧНЫЕ ВИДЫ МЫШЛЕНИЯ",
}


BLOCK_DESCRIPTIONS = {
    "base": (
        "Начало маршрута. Пользователь смотрит на карту мышления: "
        "мотивация, самооценка, амбиции и масштаб целей."
    ),
    "present": (
        "Движение по улицам настоящего. Здесь видно, как человек действует, "
        "распределяет время, ресурсы и принимает текущие решения."
    ),
    "past": (
        "Остановка у зданий прошлого опыта. Этот блок показывает, как пользователь "
        "анализирует события, причины, последствия и выводы."
    ),
    "future": (
        "Выезд на дорогу будущего. Здесь проявляются стратегия, прогнозирование, "
        "творчество и способность строить сценарии."
    ),
    "people": (
        "Встреча с людьми в городе. Блок показывает коммуникацию, эмоции, лидерство "
        "и способность действовать вместе с другими."
    ),
    "project": (
        "Рабочее пространство: офис, мастерская, проектная комната. Здесь видно, "
        "как человек превращает идеи в результат."
    ),
    "unusual": (
        "Необычные улицы города. Этот блок показывает нестандартность, парадоксы, "
        "удачу, контекст и решения в неопределённости."
    ),
}


def _safe_text(value) -> str:
    if value is None:
        return ""

    text = str(value)

    replacements = {
        "✅": "+",
        "⬜": "-",
        "📄": "",
        "📊": "",
        "📽": "",
        "🧠": "",
        "🔵": "",
        "🟢": "",
        "⚫": "",
        "🌸": "",
        "🤍": "",
        "🟡": "",
        "🟣": "",
        "🚗": "",
        "🏁": "",
        "💳": "",
        "🤖": "",
        "❓": "",
        "➡️": "",
        "⬅️": "",
        "🎤": "",
        "⏳": "",
        "🔥": "",
        "🎯": "",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    return text.strip()


def _shorten(text: str, limit: int = 650) -> str:
    text = _safe_text(text)

    if len(text) <= limit:
        return text

    return text[:limit].rstrip() + "..."


def _calculate_summary(results: dict) -> dict:
    total = len(results)

    found = sum(
        1
        for value in results.values()
        if value.get("presence") == "ЕСТЬ"
    )

    missing = total - found

    average_score = 0

    if total:
        average_score = round(
            sum(value.get("score", 0) for value in results.values()) / total,
            1,
        )

    return {
        "total": total,
        "found": found,
        "missing": missing,
        "average_score": average_score,
    }


def _found_types(results: dict) -> list[str]:
    return [
        type_name
        for type_name, data in results.items()
        if data.get("presence") == "ЕСТЬ"
    ]


def _missing_types(results: dict) -> list[str]:
    return [
        type_name
        for type_name, data in results.items()
        if data.get("presence") != "ЕСТЬ"
    ]


def _add_title(slide, title: str):
    title_box = slide.shapes.add_textbox(
        Inches(0.7),
        Inches(0.35),
        Inches(8.8),
        Inches(0.7),
    )

    frame = title_box.text_frame
    frame.clear()

    p = frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(35, 45, 70)


def _add_subtitle(slide, text: str):
    box = slide.shapes.add_textbox(
        Inches(0.75),
        Inches(1.15),
        Inches(8.5),
        Inches(0.6),
    )

    frame = box.text_frame
    frame.clear()

    p = frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(80, 90, 110)


def _add_body_text(
    slide,
    text: str,
    left: float = 0.8,
    top: float = 1.7,
    width: float = 8.6,
    height: float = 4.8,
    font_size: int = 16,
):
    box = slide.shapes.add_textbox(
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )

    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()

    for index, line in enumerate(_safe_text(text).split("\n")):
        if index == 0:
            p = frame.paragraphs[0]
        else:
            p = frame.add_paragraph()

        p.text = line
        p.font.size = Pt(font_size)
        p.font.color.rgb = RGBColor(45, 50, 65)
        p.space_after = Pt(5)

    return box


def _add_footer(slide, page_text: str = "SPIKA | Диагностика типов мышления"):
    box = slide.shapes.add_textbox(
        Inches(0.7),
        Inches(6.95),
        Inches(8.8),
        Inches(0.3),
    )

    frame = box.text_frame
    frame.clear()

    p = frame.paragraphs[0]
    p.text = page_text
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(120, 120, 120)
    p.alignment = PP_ALIGN.CENTER


def _add_stat_box(slide, x, y, title, value):
    shape = slide.shapes.add_shape(
        1,
        Inches(x),
        Inches(y),
        Inches(2.5),
        Inches(1.05),
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(245, 247, 250)
    shape.line.color.rgb = RGBColor(210, 215, 225)

    text_frame = shape.text_frame
    text_frame.clear()

    p1 = text_frame.paragraphs[0]
    p1.text = str(value)
    p1.font.size = Pt(24)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(35, 45, 70)
    p1.alignment = PP_ALIGN.CENTER

    p2 = text_frame.add_paragraph()
    p2.text = title
    p2.font.size = Pt(10)
    p2.font.color.rgb = RGBColor(80, 90, 110)
    p2.alignment = PP_ALIGN.CENTER


def _blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _group_answers_by_block(answers: list[dict]) -> dict:
    grouped = {
        block_id: []
        for block_id in BLOCK_ORDER
    }

    for item in answers:
        block_id = item.get("block_id")

        if block_id in grouped:
            grouped[block_id].append(item)

    return grouped


def _build_block_results_from_answers(answers: list[dict]) -> dict:
    grouped = {
        block_id: []
        for block_id in BLOCK_ORDER
    }

    for item in answers:
        block_id = item.get("block_id")

        if block_id in grouped:
            grouped[block_id].append(item)

    return grouped


def build_ppt_report(user_id: int, results: dict, answers: list[dict]) -> str:
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)

    file_path = REPORTS_DIR / f"report_{user_id}.pptx"

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    summary = _calculate_summary(results)
    found = _found_types(results)
    missing = _missing_types(results)

    # SLIDE 1 — TITLE
    slide = _blank_slide(prs)

    _add_title(slide, "Путешествие по Городу Мышления")

    _add_subtitle(
        slide,
        "Итоговая презентация по диагностике типов мышления",
    )

    _add_body_text(
        slide,
        (
            "Этот файл показывает маршрут прохождения опроса: "
            "от базовых целей и мотивации до нестандартных видов мышления.\n\n"
            "Пользователь проходит город своего мышления, встречает разные "
            "ситуации, анализирует опыт, принимает решения и получает карту "
            "сильных сторон и зон развития."
        ),
        top=2.0,
        font_size=18,
    )

    _add_body_text(
        slide,
        f"Дата создания: {datetime.now().strftime('%d.%m.%Y %H:%M')}",
        top=6.25,
        height=0.4,
        font_size=12,
    )

    _add_footer(slide)

    # SLIDE 2 — ROUTE
    slide = _blank_slide(prs)

    _add_title(slide, "Маршрут диагностики")

    _add_body_text(
        slide,
        (
            "1. Выход на улицу мышления: знакомство с опросом.\n"
            "2. Посадка в машину: выбор направления и старт маршрута.\n"
            "3. Движение по районам города: блоки вопросов.\n"
            "4. Встречи и остановки: анализ опыта, людей, проектов и решений.\n"
            "5. Заход в помещение: получение результатов и диагностики.\n"
            "6. Возвращение на улицу: новый взгляд на себя и своё мышление."
        ),
        top=1.55,
        font_size=17,
    )

    _add_footer(slide)

    # SLIDE 3 — SUMMARY
    slide = _blank_slide(prs)

    _add_title(slide, "Сводка результата")

    _add_stat_box(slide, 0.85, 1.55, "Проверено типов", summary["total"])
    _add_stat_box(slide, 3.75, 1.55, "Найдено", summary["found"])
    _add_stat_box(slide, 6.65, 1.55, "Зоны развития", summary["missing"])

    _add_body_text(
        slide,
        f"Средний балл: {summary['average_score']}/10",
        top=3.0,
        font_size=20,
    )

    if summary["total"] == 0:
        diagnostic = (
            "Диагностика пока не сформирована. Нужно пройти вопросы, "
            "чтобы получить карту типов мышления."
        )
    elif summary["found"] >= max(1, summary["total"] * 0.7):
        diagnostic = (
            "Ответы показывают широкий спектр проявленных типов мышления. "
            "Пользователь уверенно раскрывает логику, опыт, решения и выводы."
        )
    elif summary["found"] >= max(1, summary["total"] * 0.4):
        diagnostic = (
            "Есть выраженная база мышления и несколько сильных зон. "
            "Часть направлений требует дополнительной конкретики и практики."
        )
    else:
        diagnostic = (
            "Пока раскрыта только часть типов мышления. Для лучшего результата "
            "нужно больше примеров, действий, решений и личных выводов."
        )

    _add_body_text(
        slide,
        diagnostic,
        top=3.7,
        font_size=16,
    )

    _add_footer(slide)

    # SLIDE 4 — BLOCK RESULTS
    block_results = _build_block_results_from_answers(answers)

    for block_id in BLOCK_ORDER:
        block_items = block_results.get(block_id, [])

        if not block_items:
            continue

        slide = _blank_slide(prs)

        _add_title(slide, BLOCK_TITLES.get(block_id, block_id))

        _add_subtitle(
            slide,
            BLOCK_DESCRIPTIONS.get(block_id, ""),
        )

        lines = []

        for item in block_items:
            type_name = item.get("type", "")
            score = item.get("score", 0)
            presence = item.get("presence", "НЕТ")
            marker = "+" if presence == "ЕСТЬ" else "-"
            lines.append(f"{marker} {type_name} — {score}/10")

        text = "\n".join(lines[:12])

        _add_body_text(
            slide,
            text,
            top=2.0,
            font_size=13,
        )

        _add_footer(slide)

    # SLIDE — STRONG SIDES
    slide = _blank_slide(prs)

    _add_title(slide, "Сильные стороны")

    if found:
        found_text = "\n".join(
            f"+ {type_name} — {results[type_name].get('score', 0)}/10"
            for type_name in found[:18]
        )
    else:
        found_text = (
            "Сильные стороны пока не выявлены. "
            "Нужно пройти больше вопросов или раскрывать ответы подробнее."
        )

    _add_body_text(
        slide,
        found_text,
        top=1.55,
        font_size=14,
    )

    _add_footer(slide)

    # SLIDE — DEVELOPMENT ZONES
    slide = _blank_slide(prs)

    _add_title(slide, "Зоны развития")

    if missing:
        missing_text = "\n".join(
            f"- {type_name} — {results[type_name].get('score', 0)}/10"
            for type_name in missing[:18]
        )
    else:
        missing_text = (
            "Все проверенные типы мышления проявлены. "
            "Дальше важно развивать глубину, качество и применение мышления."
        )

    _add_body_text(
        slide,
        missing_text,
        top=1.55,
        font_size=14,
    )

    _add_footer(slide)

           # ANSWER ANALYSIS SLIDES
    if answers:
        for index, item in enumerate(answers, start=1):
            thinking_type = _safe_text(item.get("type", ""))
            question = _safe_text(item.get("question", ""))
            answer = _safe_text(item.get("answer", ""))

            short_analysis = _safe_text(
                item.get("analysis", "") or "Краткий анализ не сформирован."
            )

            full_analysis = _safe_text(
                item.get("full_analysis", "") or "Расширенный анализ не сформирован."
            )

            values_analysis = _safe_text(
                item.get("values_analysis", "") or "Ценностная диагностика не сформирована."
            )

            detected_values = _safe_text(
                item.get("detected_values", "") or "явно не выявлены"
            )

            detected_desires = _safe_text(
                item.get("detected_desires", "") or "явно не выявлены"
            )

            detected_importance = _safe_text(
                item.get("detected_importance", "") or "явно не выявлены"
            )

            contradictions = _safe_text(
                item.get("contradictions", "") or "данных недостаточно"
            )

            responsibility = _safe_text(
                item.get("responsibility", "") or "данных недостаточно"
            )

            responsibility_shift = _safe_text(
                item.get("responsibility_shift", "") or "явно не выявлено"
            )

            advice = _safe_text(
                item.get("advice", "") or "Совет не сформирован."
            )

            score = item.get("score", 0)
            presence = item.get("presence", "НЕТ")

            # Slide 1: thinking analysis
            slide = _blank_slide(prs)

            _add_title(slide, f"Анализ ответа {index}")

            _add_subtitle(
                slide,
                f"{thinking_type} | {score}/10 | {presence}",
            )

            body = (
                f"Вопрос:\n"
                f"{_shorten(question, 220)}\n\n"
                f"Ответ пользователя:\n"
                f"{_shorten(answer, 360)}\n\n"
                f"Что видно по ответу:\n"
                f"{_shorten(short_analysis, 360)}\n\n"
                f"Подробная диагностика:\n"
                f"{_shorten(full_analysis, 720)}\n\n"
                f"Ориентир дальше:\n"
                f"{_shorten(advice, 260)}"
            )

            _add_body_text(
                slide,
                body,
                top=1.55,
                height=5.25,
                font_size=10,
            )

            _add_footer(slide)

            # Slide 2: values analysis
            slide = _blank_slide(prs)

            _add_title(slide, f"Ценности ответа {index}")

            _add_subtitle(
                slide,
                f"{thinking_type} | ценностные сигналы",
            )

            body = (
                f"Ценностная диагностика:\n"
                f"{_shorten(values_analysis, 650)}\n\n"
                f"Выявленные ценности:\n"
                f"{_shorten(detected_values, 260)}\n\n"
                f"Выявленные желания:\n"
                f"{_shorten(detected_desires, 260)}\n\n"
                f"Выявленные важности:\n"
                f"{_shorten(detected_importance, 260)}\n\n"
                f"Возможные противоречия:\n"
                f"{_shorten(contradictions, 300)}\n\n"
                f"Ответственность:\n"
                f"{_shorten(responsibility, 280)}\n\n"
                f"Перекладывание ответственности:\n"
                f"{_shorten(responsibility_shift, 240)}\n\n"
            )

            _add_body_text(
                slide,
                body,
                top=1.55,
                height=5.25,
                font_size=9,
            )

            _add_footer(slide)
    
    # RECOMMENDATIONS
    slide = _blank_slide(prs)

    _add_title(slide, "Рекомендации")

    recommendations = (
        "1. Отвечать через реальные ситуации: что произошло, что было сделано, какой вывод появился.\n"
        "2. Для зон развития выбрать 1-2 типа мышления и тренировать их через практические задачи.\n"
        "3. Использовать сильные стороны как опору для проектов, общения, решений и планирования.\n"
        "4. Возвращаться к диагностике после обучения или практики, чтобы увидеть динамику.\n"
        "5. Смотреть на отчёт как на карту маршрута: сильные зоны показывают опору, зоны развития — направление движения."
    )

    _add_body_text(
        slide,
        recommendations,
        top=1.55,
        font_size=15,
    )

    _add_footer(slide)

    # PROJECT DEVELOPMENT
    slide = _blank_slide(prs)

    _add_title(slide, "Блок дальнейшего развития проекта")

    _add_body_text(
        slide,
        (
            "Здесь будут ссылки на курсы по обучению.\n"
            "Здесь будут ссылки для оплаты обучения.\n"
            "Здесь будет информация по оплате опроса.\n"
            "Здесь будет контактная информация для связи с экспертом.\n"
            "Здесь будет информация для записи на консультацию."
        ),
        top=1.55,
        font_size=16,
    )

    _add_footer(slide)

    # FINAL SLIDE
    slide = _blank_slide(prs)

    _add_title(slide, "Финальная сцена маршрута")

    _add_body_text(
        slide,
        (
            "Пользователь выходит из помещения с результатами диагностики "
            "и возвращается на улицу с новым взглядом на себя.\n\n"
            "Теперь мышление воспринимается не как абстрактная способность, "
            "а как инструмент, который помогает двигаться по реальному миру: "
            "принимать решения, общаться, строить планы, запускать проекты "
            "и понимать собственные сильные стороны."
        ),
        top=1.75,
        font_size=18,
    )

    _add_footer(slide, "Конец маршрута — начало нового понимания")

    prs.save(file_path)

    return str(file_path)